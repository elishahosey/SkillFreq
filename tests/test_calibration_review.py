import csv
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from skillfreq.calibration_review import (
    SUBGRADE_OPTIONS, build_queue, default_review_path, find_review, identity, load_comparison, load_reviews, save_review,
)


class CalibrationReviewTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.source = self.directory / 'comparison.csv'
        self.rows = [dict(id='001', source_site='site', title='Changed decision', company='Example',
                         description='Build pipelines', old_lane='target', new_lane='target',
                         old_score=80, new_score=80, old_decision='skip', new_decision='apply_now'),
                     dict(id='002', source_site='site', title='Changed score', company='', description='',
                          old_lane='target', new_lane='target', old_score=50, new_score=65,
                          old_decision='skip', new_decision='skip'),
                     dict(id='003', source_site='site', title='Unchanged', company='', description='',
                          old_lane='target', new_lane='target', old_score=50, new_score=50,
                          old_decision='skip', new_decision='skip')]
        with self.source.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=self.rows[0])
            writer.writeheader()
            writer.writerows(self.rows)
        self.records, _ = load_comparison(self.source)
        self.output = self.directory / 'calibration_review_2026-09-25.xlsx'

    def test_save_update_and_resume_preserve_source_and_literal_notes(self):
        original = self.source.read_bytes()
        row = self.records[0]
        save_review(self.output, row, 1, decision_quality='better', score_quality='reasonable',
                    requirement_interpretation='correct')
        save_review(self.output, row, 3, '=literal note', ['stack_mismatch'],
                    decision_quality='worse', score_quality='too_high', requirement_interpretation='incorrect')
        saved = load_reviews(self.output, row)
        self.assertEqual(len(saved), 1)
        review = saved[identity(row)]
        self.assertEqual(review['job_id'], '001')
        self.assertEqual(review['review_rating'], 3)
        self.assertEqual(review['new_fit_score'], 80)
        self.assertEqual(review['review_note'], '=literal note')
        self.assertEqual(review['issue_tags'], 'stack_mismatch')
        self.assertEqual(review['decision_quality'], 'worse')
        self.assertEqual(review['score_quality'], 'too_high')
        self.assertEqual(review['requirement_interpretation'], 'incorrect')
        self.assertTrue(review['reviewed'])
        self.assertTrue(review['reviewed_at'])
        self.assertEqual(default_review_path(row), self.output)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual([r['job_id'] for r in build_queue(self.records, saved, unchanged_sample=0)], ['002'])

    def test_queue_filters_priorities_thresholds_and_sampling(self):
        self.assertEqual([r['job_id'] for r in build_queue(self.records, {}, unchanged_sample=0)], ['001', '002'])
        self.assertEqual([r['job_id'] for r in build_queue(self.records, {}, threshold=20, unchanged_sample=0)], ['001'])
        self.assertEqual([r['job_id'] for r in build_queue(self.records, {}, signals=[], unchanged_sample=5)], ['003'])
        reviewed = {identity(self.records[0]): {}}
        self.assertEqual([r['job_id'] for r in build_queue(self.records, reviewed, status='Reviewed')], ['001'])
        self.assertEqual([r['job_id'] for r in build_queue(self.records, reviewed, status='All', unchanged_sample=0)], ['002', '001'])
        self.records[2]['new_role_lane'] = 'secondary'
        self.assertEqual([r['job_id'] for r in build_queue(self.records, {}, signals=['Lane changed'], unchanged_sample=0)], ['003'])
        self.records[2]['new_ai_review_required'] = 'True'
        self.assertEqual([r['job_id'] for r in build_queue(self.records, {}, signals=['AI review changed'], unchanged_sample=0)], ['003'])

    def test_excel_named_fields_and_missing_context(self):
        path = self.directory / 'input.xlsx'
        book = Workbook()
        book.active.append(['new_fit_score', 'source_row_id', 'old_fit_score'])
        book.active.append([90, '0007', 50])
        book.save(path)
        book.close()
        original = path.read_bytes()
        rows, _ = load_comparison(path)
        self.assertEqual(rows[0]['job_id'], '0007')
        self.assertEqual(rows[0]['company'], '')
        with self.assertRaisesRegex(ValueError, 'separate'):
            save_review(path, rows[0], 1)
        self.assertEqual(path.read_bytes(), original)

    def test_total_target_counts_saved_reviews_across_filters_and_rebuilds(self):
        rows = [dict(self.records[0], job_id=str(i)) for i in range(120)]
        reviews = {identity(row): {} for row in rows[:6]}
        queue = build_queue(rows, reviews, limit=100, unchanged_sample=0)
        self.assertEqual(len(queue), 94)
        self.assertFalse(any(identity(row) in reviews for row in queue))
        reviews[identity(queue[0])] = {}
        self.assertEqual(len(build_queue(rows, reviews, limit=100, unchanged_sample=0)), 93)
        # Saved reviews count even when they no longer meet the change filters.
        for row in rows[:6]:
            row['new_apply_decision'] = row['old_apply_decision']
        reviews[('unrelated', 'job')] = {}
        self.assertEqual(len(build_queue(rows, reviews, limit=100, unchanged_sample=0)), 93)
        self.assertEqual(len(build_queue(rows, reviews, limit=100, unchanged_sample=0, status='All')), 100)
        self.assertEqual(len(build_queue(rows, reviews, limit=100, status='Reviewed')), 7)
        self.assertEqual(build_queue(rows, reviews, limit=7), [])
        self.assertEqual(build_queue(rows, reviews, limit=3), [])
        self.assertEqual(len(build_queue(rows, reviews, limit=3, status='Reviewed')), 7)
        self.assertEqual(len(build_queue(rows, reviews, limit=10, unchanged_sample=0)), 3)

    def test_unchanged_sample_is_part_of_total_and_counts_saved_sample(self):
        changed = [dict(self.records[0], job_id=str(i)) for i in range(10)]
        unchanged = [dict(self.records[2], job_id=str(i)) for i in range(10, 20)]
        rows = changed + unchanged
        reviews = {identity(changed[0]): {}, identity(unchanged[0]): {}}
        queue = build_queue(rows, reviews, limit=6, unchanged_sample=3)
        self.assertEqual(len(queue), 4)
        self.assertEqual(sum(row in unchanged for row in queue), 2)
        self.assertEqual(queue, build_queue(rows, reviews, limit=6, unchanged_sample=3))
        for row in queue:
            reviews[identity(row)] = {}
        self.assertEqual(build_queue(rows, reviews, limit=6, unchanged_sample=3), [])
        more = build_queue(rows, reviews, limit=8, unchanged_sample=3)
        self.assertEqual(len(more), 2)
        self.assertTrue(all(row in changed for row in more))
        self.assertEqual(len(build_queue(rows, {}, limit=2, unchanged_sample=5)), 2)

    def test_two_exports_match_source_identity_and_read_grade_evidence(self):
        import json
        paths = [self.directory / 'old.csv', self.directory / 'new.csv']
        for path, lane in zip(paths, ['target', 'secondary']):
            with path.open('w', encoding='utf-8', newline='') as stream:
                writer = csv.DictWriter(stream, fieldnames=['id', 'source_site', 'title', 'grade_json'])
                writer.writeheader()
                writer.writerow(dict(id='001', source_site='site', title='Engineer',
                                     grade_json=json.dumps(dict(role_lane=lane, fit_score=50,
                                         blocking_reasons=['experience'], grading_version=lane))))
        rows, _ = load_comparison(paths[1], paths[0])
        self.assertEqual(rows[0]['old_role_lane'], 'target')
        self.assertEqual(rows[0]['new_role_lane'], 'secondary')
        self.assertIn('experience', rows[0]['new_blocking_reasons'])
        self.assertEqual(rows[0]['new_grading_version'], 'secondary')

    def test_duplicate_input_and_unrelated_output_rejected(self):
        with self.source.open('a', encoding='utf-8', newline='') as stream:
            csv.DictWriter(stream, fieldnames=self.rows[0]).writerow(self.rows[0])
        with self.assertRaisesRegex(ValueError, 'Duplicate'):
            load_comparison(self.source)
        save_review(self.output, self.records[0], 1)
        other = dict(self.records[0], comparison_after='different.csv')
        with self.assertRaisesRegex(ValueError, 'different comparison'):
            save_review(self.output, other, 2)
        self.assertEqual(load_reviews(self.output)[identity(self.records[0])]['review_rating'], 1)

    def test_failed_replace_keeps_saved_review(self):
        save_review(self.output, self.records[0], 1)
        with patch('pathlib.Path.replace', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                save_review(self.output, self.records[0], 3)
        self.assertEqual(load_reviews(self.output)[identity(self.records[0])]['review_rating'], 1)
        self.assertEqual(list(self.directory.glob('*.xlsx')), [self.output])

    def test_find_review_by_exact_id_and_reject_unknown_or_ambiguous_id(self):
        save_review(self.output, self.records[0], 3, 'Check requirements')
        reviews = load_reviews(self.output)
        self.assertEqual(find_review(reviews, ' 001 ')['review_note'], 'Check requirements')
        for job_id in ('', '1', 'unknown'):
            with self.assertRaises(LookupError):
                find_review(reviews, job_id)
        save_review(self.output, dict(self.records[0], source_site='another_site'), 1)
        with self.assertRaisesRegex(ValueError, 'Duplicate review rows'):
            find_review(load_reviews(self.output), '001')
        original = self.output.read_bytes()
        with self.assertRaisesRegex(ValueError, 'Duplicate review rows'):
            save_review(self.output, self.records[0], 2, require_existing=True)
        self.assertEqual(self.output.read_bytes(), original)

    def test_duplicate_review_rows_and_missing_edit_target_never_write(self):
        with self.assertRaisesRegex(ValueError, 'No saved review'):
            save_review(self.output, self.records[0], 1, require_existing=True)
        self.assertFalse(self.output.exists())
        save_review(self.output, self.records[0], 1)
        book = load_workbook(self.output)
        book.active.append([cell.value for cell in book.active[2]])
        book.save(self.output)
        book.close()
        original = self.output.read_bytes()
        with self.assertRaisesRegex(ValueError, 'Duplicate review rows'):
            load_reviews(self.output)
        with self.assertRaisesRegex(ValueError, 'Duplicate review rows'):
            save_review(self.output, self.records[0], 3, require_existing=True)
        self.assertEqual(self.output.read_bytes(), original)

    def test_edit_preserves_other_records_extra_columns_and_formulas(self):
        # Save in a different order from the source to verify identity-based updates.
        save_review(self.output, self.records[1], 2, 'Keep this note', ['posting_ambiguity'])
        save_review(self.output, self.records[0], 1)
        book = load_workbook(self.output)
        sheet = book.active
        extra = sheet.max_column + 1
        sheet.cell(1, extra, 'custom_metadata')
        extra_letter = sheet.cell(1, extra).column_letter
        sheet.column_dimensions[extra_letter].width = 42
        sheet.cell(2, extra, 'Other record metadata')
        sheet.cell(3, extra, 'Edited record metadata')
        sheet.cell(1, extra + 1, 'custom_formula')
        sheet.cell(3, extra + 1, '=1+2')
        book.create_sheet('Extra').cell(1, 1, 'Keep this sheet')
        book.save(self.output)
        original_other = [cell.value for cell in sheet[2]]
        original_headers = [cell.value for cell in sheet[1]]
        book.close()
        save_review(self.output, self.records[0], 3, 'Corrected', ['stack_mismatch'],
                    decision_quality='worse', score_quality='too_high',
                    requirement_interpretation='incorrect', require_existing=True)
        reviews = load_reviews(self.output)
        self.assertEqual(len(reviews), 2)
        self.assertEqual(find_review(reviews, '001')['review_rating'], 3)
        self.assertEqual(find_review(reviews, '001')['review_note'], 'Corrected')
        book = load_workbook(self.output)
        self.assertEqual(book.active.max_row, 3)
        self.assertEqual([cell.value for cell in book.active[1]], original_headers)
        self.assertEqual([cell.value for cell in book.active[2]], original_other)
        self.assertEqual(book.active.cell(3, extra).value, 'Edited record metadata')
        self.assertEqual(book.active.column_dimensions[extra_letter].width, 42)
        self.assertEqual(book.active.cell(3, extra + 1).value, '=1+2')
        self.assertEqual(book['Extra'].cell(1, 1).value, 'Keep this sheet')
        book.close()

    def test_legacy_workbook_loads_and_gains_subgrades_without_losing_reviews(self):
        # An actual old-format workbook: no sub-grade columns, not even blank ones.
        legacy = dict(self.records[0], review_rating=2, review_note='Original note',
                      issue_tags='posting_ambiguity', reviewed=True, reviewed_at='2026-09-25T12:00:00+00:00')
        book = Workbook()
        book.active.append(list(legacy))
        book.active.append(list(legacy.values()))
        book.save(self.output)
        book.close()
        reviews = load_reviews(self.output)
        self.assertTrue(all(reviews[identity(self.records[0])][f] == '' for f in SUBGRADE_OPTIONS))
        app = self.app()
        next(w for w in app.selectbox if w.label == 'Review status').select('Reviewed')
        self.click(app, 'Load / rebuild queue')
        self.assertEqual(app.radio[0].value, 2)
        self.assertEqual(app.text_area[0].value, 'Original note')
        self.assertEqual(next(w for w in app.multiselect if w.label == 'Issue tags (optional)').value,
                         ['posting_ambiguity'])
        for label in ('Decision quality', 'Score quality', 'Requirement interpretation'):
            self.assertIsNone(next(w for w in app.selectbox if w.label == label).value)
        save_review(self.output, self.records[1], 1, decision_quality='same', score_quality='reasonable',
                    requirement_interpretation='correct')
        self.click(app, 'Save')
        reviews = load_reviews(self.output)
        self.assertEqual(len(reviews), 2)
        self.assertEqual(reviews[identity(self.records[0])]['review_note'], 'Original note')
        self.assertEqual(reviews[identity(self.records[0])]['issue_tags'], 'posting_ambiguity')
        self.assertEqual(reviews[identity(self.records[1])]['decision_quality'], 'same')
        self.assertEqual(reviews[identity(self.records[1])]['score_quality'], 'reasonable')
        self.assertEqual(reviews[identity(self.records[1])]['requirement_interpretation'], 'correct')

    def app(self):
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / 'scripts/review_calibration.py',
                                default_timeout=15).run()
        app.selectbox[0].select('One comparison file')
        app.text_input[1].input(str(self.source))
        app.text_input[2].input('')  # Discover dated progress on a new session.
        app.button[0].click().run()
        self.assertFalse(app.exception)
        return app

    def click(self, app, label):
        next(button for button in app.button if button.label == label).click().run()
        self.assertFalse(app.exception)

    def test_ui_save_navigate_update_and_restart(self):
        app = self.app()
        self.assertEqual(app.subheader[0].value, 'Changed decision')
        self.click(app, 'Save + Next')
        self.assertTrue(app.error)
        self.assertEqual(app.subheader[0].value, 'Changed decision')
        self.assertFalse(list(self.directory.glob('calibration_review_*.xlsx')))
        labels = ('Decision quality', 'Score quality', 'Requirement interpretation')
        first_grades = ('better', 'reasonable', 'correct')
        for label, value in zip(labels, first_grades):
            widget = next(w for w in app.selectbox if w.label == label)
            self.assertIsNone(widget.value)
            widget.select(value)
        next(w for w in app.multiselect if w.label == 'Issue tags (optional)').select('requirement_semantics_issue')
        app.radio[0].set_value(1)
        self.click(app, 'Save + Next')
        self.assertEqual(app.subheader[0].value, 'Changed score')
        self.assertIsNone(app.radio[0].value)
        for label in labels:
            self.assertIsNone(next(w for w in app.selectbox if w.label == label).value)
        app.text_area[0].input('Unsaved draft')
        self.click(app, 'Previous')
        self.assertEqual(app.radio[0].value, 1)
        for label, value in zip(labels, first_grades):
            self.assertEqual(next(w for w in app.selectbox if w.label == label).value, value)
        self.assertEqual(next(w for w in app.multiselect if w.label == 'Issue tags (optional)').value,
                         ['requirement_semantics_issue'])
        updated_grades = ('worse', 'too_low', 'incorrect')
        for label, value in zip(labels, updated_grades):
            next(w for w in app.selectbox if w.label == label).select(value)
        app.radio[0].set_value(3)
        app.text_area[0].input('Decision needs correction')
        self.click(app, 'Save')
        path = default_review_path(self.records[0])
        self.assertEqual(len(load_reviews(path)), 1)
        self.assertEqual(load_reviews(path)[identity(self.records[0])]['review_rating'], 3)
        restarted = self.app()
        self.assertEqual(restarted.subheader[0].value, 'Changed score')
        self.assertEqual(restarted.text_area[0].value, '')
        restarted.radio[0].set_value(2)
        self.click(restarted, 'Save + Next')
        self.assertTrue(any('Consider adding' in info.value for info in restarted.info))
        self.assertEqual(len(load_reviews(path)), 2)
        reopened = self.app()
        next(w for w in reopened.selectbox if w.label == 'Review status').select('Reviewed')
        self.click(reopened, 'Load / rebuild queue')
        self.assertEqual(reopened.subheader[0].value, 'Changed decision')
        self.assertEqual(reopened.radio[0].value, 3)
        self.assertEqual(reopened.text_area[0].value, 'Decision needs correction')
        for label, value in zip(labels, updated_grades):
            self.assertEqual(next(w for w in reopened.selectbox if w.label == label).value, value)
        self.assertEqual(next(w for w in reopened.multiselect if w.label == 'Issue tags (optional)').value,
                         ['requirement_semantics_issue'])

    def test_ui_cumulative_progress_completion_and_target_changes(self):
        save_review(self.output, self.records[0], 1)
        app = self.app()
        next(w for w in app.number_input if w.label == 'Total review target').set_value(2)
        next(w for w in app.number_input if w.label == 'Unchanged sanity sample').set_value(0)
        self.click(app, 'Load / rebuild queue')
        self.assertEqual(app.get('progress')[0].proto.text, '1 / 2 total records reviewed')
        self.assertEqual(len(app.session_state.run['queue']), 1)
        app.radio[0].set_value(1)
        self.click(app, 'Save + Next')
        self.assertEqual(app.get('progress')[0].proto.text, '2 / 2 total records reviewed')
        self.assertTrue(any('target reached' in item.value for item in app.success))
        self.click(app, 'Load / rebuild queue')
        self.assertEqual(app.get('progress')[0].proto.text, '2 / 2 total records reviewed')
        self.assertFalse(app.subheader)
        next(w for w in app.number_input if w.label == 'Total review target').set_value(1)
        self.click(app, 'Load / rebuild queue')
        self.assertEqual(app.get('progress')[0].proto.text, '2 / 1 total records reviewed')
        self.assertEqual(app.get('progress')[0].proto.value, 100)
        next(w for w in app.selectbox if w.label == 'Review status').select('Reviewed')
        self.click(app, 'Load / rebuild queue')
        self.assertEqual(len(app.session_state.run['queue']), 2)
        self.assertEqual(app.radio[0].value, 1)
        restarted = self.app()
        self.assertEqual(restarted.get('progress')[0].proto.text, '2 / 100 total records reviewed')
        self.assertEqual(restarted.subheader[0].value, 'Unchanged')
        self.assertTrue(any('fewer records' in item.value for item in restarted.info))

    def test_ui_edit_existing_review_and_return_to_queue(self):
        save_review(self.output, self.records[0], 3, 'Original note', ['stack_mismatch', 'legacy_tag'],
                    decision_quality='worse', score_quality='too_high', requirement_interpretation='incorrect')
        app = self.app()
        self.click(app, 'Next')
        original_queue = app.session_state.run['queue']
        original_index = app.session_state.run['index']
        next(w for w in app.text_input if w.label == 'Job ID').input('001')
        self.click(app, 'Load review')
        self.assertEqual(app.subheader[0].value, 'Changed decision')
        self.assertTrue(any('Editing existing review: 001' in item.value for item in app.info))
        self.assertTrue(any(item.value == 'Build pipelines' for item in app.text))
        self.assertTrue(app.table)
        self.assertEqual(app.radio[0].value, 3)
        self.assertEqual(app.text_area[0].value, 'Original note')
        self.assertEqual(next(w for w in app.multiselect if w.label == 'Issue tags (optional)').value,
                         ['stack_mismatch', 'legacy_tag'])
        labels = ('Decision quality', 'Score quality', 'Requirement interpretation')
        for label, expected in zip(labels, ('worse', 'too_high', 'incorrect')):
            self.assertEqual(next(w for w in app.selectbox if w.label == label).value, expected)
        # Loading the same ID again restores saved answers rather than stale drafts.
        app.text_area[0].input('Discard this draft')
        self.click(app, 'Load review')
        self.assertEqual(app.text_area[0].value, 'Original note')
        app.radio[0].set_value(1)
        app.text_area[0].input('Corrected note')
        for label, value in zip(labels, ('better', 'reasonable', 'correct')):
            next(w for w in app.selectbox if w.label == label).select(value)
        next(w for w in app.multiselect if w.label == 'Issue tags (optional)').unselect('stack_mismatch')
        self.click(app, 'Save')
        reviews = load_reviews(self.output)
        self.assertEqual(len(reviews), 1)
        edited = find_review(reviews, '001')
        self.assertEqual(edited['review_rating'], 1)
        self.assertEqual(edited['review_note'], 'Corrected note')
        self.assertEqual(edited['issue_tags'], 'legacy_tag')
        for field, value in zip(SUBGRADE_OPTIONS, ('better', 'reasonable', 'correct')):
            self.assertEqual(edited[field], value)
        self.assertEqual(app.get('progress')[0].proto.text, '1 / 100 total records reviewed')
        self.click(app, 'Return to review queue')
        self.assertEqual(app.session_state.run['queue'], original_queue)
        self.assertEqual(app.session_state.run['index'], original_index)
        self.assertEqual(app.subheader[0].value, 'Unchanged')
        self.assertIsNone(app.radio[0].value)
        app.radio[0].set_value(1)
        self.click(app, 'Save')
        self.assertEqual(len(load_reviews(self.output)), 2)

    def test_ui_edit_lookup_errors_and_completed_queue(self):
        save_review(self.output, self.records[0], 1)
        app = self.app()
        next(w for w in app.number_input if w.label == 'Total review target').set_value(1)
        self.click(app, 'Load / rebuild queue')
        self.assertFalse(app.subheader)
        next(w for w in app.text_input if w.label == 'Job ID').input('unknown')
        self.click(app, 'Load review')
        self.assertTrue(any('No saved review found' in item.value for item in app.warning))
        next(w for w in app.text_input if w.label == 'Job ID').input('001')
        self.click(app, 'Load review')
        self.assertEqual(app.radio[0].value, 1)
        self.click(app, 'Return to review queue')
        self.assertFalse(app.subheader)
        # A review added to disk after loading must be read on lookup.
        save_review(self.output, dict(self.records[0], job_id='missing-source'), 1)
        next(w for w in app.text_input if w.label == 'Job ID').input('missing-source')
        self.click(app, 'Load review')
        self.assertTrue(any('evidence is missing' in item.value for item in app.warning))
        save_review(self.output, dict(self.records[0], source_site='another_site'), 1)
        next(w for w in app.text_input if w.label == 'Job ID').input('001')
        self.click(app, 'Load review')
        self.assertTrue(any('Duplicate review rows' in item.value for item in app.error))
        self.assertIsNone(app.session_state.run['editing_key'])


if __name__ == '__main__':
    unittest.main()
