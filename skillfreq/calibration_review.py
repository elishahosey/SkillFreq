"""Review saved grading output without importing or running the grader."""
import csv
from datetime import date, datetime, timezone
import json
from pathlib import Path
import random
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


ISSUE_TAGS = [
    'hard_filter_problem', 'seniority_mismatch', 'skill_extraction_issue',
    'stack_mismatch', 'lane_assignment_issue', 'score_weighting_issue',
    'ai_review_disagreement', 'posting_ambiguity', 'missing_context',
    'requirement_semantics_issue', 'other',
]
SUBGRADE_OPTIONS = {
    'decision_quality': ('better', 'same', 'worse', 'unclear'),
    'score_quality': ('too_high', 'reasonable', 'too_low', 'unclear'),
    'requirement_interpretation': ('correct', 'unclear', 'incorrect'),
}
SUMMARY_FIELDS = ('role_lane', 'fit_score', 'fit_quality', 'apply_decision',
                  'pre_ai_score', 'learning_score', 'ai_review_required', 'confidence')
EVIDENCE_FIELDS = ('blocking_reasons', 'exclusion_signals', 'seniority_signals',
                   'matched_atomic_skills', 'matched_capability_concepts',
                   'missing_required_skills', 'missing_preferred_skills',
                   'review_flags', 'ai_review_reasons', 'reason_codes')
GRADE_FIELDS = SUMMARY_FIELDS + EVIDENCE_FIELDS + ('grading_version',)
SIGNALS = ('Decision changed', 'Lane changed', 'Score changed', 'AI review changed',
           'Other evidence changed')
ALIASES = {'lane': 'role_lane', 'fit': 'fit_score', 'score': 'fit_score',
           'decision': 'apply_decision', 'result': 'fit_quality'}


def text(value):
    if value is None:
        return ''
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value).strip()


def read_rows(path):
    """Read named columns, retaining text IDs and closing source handles."""
    path = Path(path)
    if path.suffix.lower() == '.csv':
        csv.field_size_limit(10_000_000)
        with path.open(encoding='utf-8-sig', newline='') as stream:
            yield from csv.DictReader(stream)
    elif path.suffix.lower() == '.xlsx':
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = book.active.iter_rows(values_only=True)
            headers = [text(v) for v in next(rows, ())]
            if len(headers) != len(set(headers)):
                raise ValueError('Duplicate Excel column names.')
            for values in rows:
                if any(v is not None for v in values):
                    yield dict(zip(headers, values))
        finally:
            book.close()
    else:
        raise ValueError('Choose a .csv or .xlsx file (first worksheet).')


def identity(row):
    # Same source identity as scripts/compare_calibration.py, plus Excel aliases.
    job_id = next((text(row.get(f)) for f in ('job_id', 'id', 'source_row_id', 'source')
                   if text(row.get(f))), '')
    if not job_id:
        raise ValueError('Every row needs id, job_id, source_row_id, or source URL.')
    return text(row.get('source_site')), job_id


def grade_fields(row):
    grade = json.loads(row['grade_json']) if row.get('grade_json') else {}
    if not isinstance(grade, dict):
        raise ValueError('grade_json must contain a JSON object.')
    result = {f: row.get(f, grade.get(f, '')) for f in GRADE_FIELDS}
    skills = result.get('matched_atomic_skills')
    if isinstance(skills, list):
        result['matched_atomic_skills'] = [s.get('canonical_skill', s) if isinstance(s, dict) else s
                                          for s in skills]
    return {f: text(v) for f, v in result.items()}


def load_comparison(after, before=None):
    """Load two grading exports, or a single old_*/new_* comparison file."""
    def load(path, combined=False):
        indexed = {}
        for raw in read_rows(path):
            key = identity(raw)
            if key in indexed:
                raise ValueError(f'Duplicate source identity in {path}: {key}')
            row = {f: text(raw.get(f)) for f in ('title', 'company', 'description', 'source')}
            row.update(source_site=key[0], job_id=key[1])
            if combined:
                for side in ('old', 'new'):
                    fields = {k[len(side)+1:]: v for k, v in raw.items() if k.startswith(side + '_')}
                    for alias, canonical in ALIASES.items():
                        if alias in fields and canonical not in fields:
                            fields[canonical] = fields[alias]
                    if not set(fields) & (set(GRADE_FIELDS) | {'grade_json'}):
                        raise ValueError('Comparison needs old_* and new_* grading fields; otherwise choose two grading exports.')
                    row.update({f'{side}_{k}': v for k, v in grade_fields(fields).items()})
            else:
                if not set(raw) & (set(GRADE_FIELDS) | {'grade_json'}):
                    raise ValueError('Input has no saved grading fields.')
                row.update(grade_fields(raw))
            indexed[key] = row
        if not indexed:
            raise ValueError(f'No records in {path}.')
        return indexed

    after = Path(after).resolve()
    new = load(after, combined=before is None)
    if before is None:
        rows = list(new.values())
        message = f'{len(rows):,} comparison records loaded.'
    else:
        before = Path(before).resolve()
        old = load(before)
        common = sorted(old.keys() & new.keys())
        if not common:
            raise ValueError('No matching source identities between OLD and NEW.')
        rows = []
        for key in common:
            a, b = old[key], new[key]
            row = {f: b[f] or a[f] for f in ('title', 'company', 'source', 'source_site', 'job_id')}
            row['description'] = b['description']
            row['old_description'] = a['description'] if a['description'] != b['description'] else ''
            for side, values in (('old', a), ('new', b)):
                row.update({f'{side}_{f}': values[f] for f in GRADE_FIELDS})
            rows.append(row)
        message = (f'{len(rows):,} matched records. Excluded {len(old)-len(common):,} OLD-only '
                   f'and {len(new)-len(common):,} NEW-only records.')
    for row in rows:
        row['comparison_before'] = str(before) if before else ''
        row['comparison_after'] = str(after)
    return rows, message


def score_delta(row):
    try:
        return abs(float(row['new_fit_score']) - float(row['old_fit_score']))
    except (ValueError, KeyError, TypeError):
        return 0.0


def changes(row, threshold=5):
    def changed(field):
        a, b = (text(row.get(side + '_' + field)).casefold() for side in ('old', 'new'))
        if field in ('fit_score', 'pre_ai_score', 'learning_score', 'confidence'):
            try:
                return float(a) != float(b)
            except ValueError:
                pass
        return a != b
    return {
        'Decision changed': changed('apply_decision'),
        'Lane changed': changed('role_lane'),
        'Score changed': changed('fit_score') and score_delta(row) >= threshold,
        'AI review changed': changed('ai_review_required') or changed('ai_review_reasons'),
        'Other evidence changed': any(changed(f) for f in EVIDENCE_FIELDS +
                                      ('fit_quality', 'pre_ai_score', 'learning_score', 'confidence', 'fit_score')),
    }


def build_queue(rows, reviews, signals=SIGNALS[:4], threshold=5, limit=100,
                unchanged_sample=5, status='Unreviewed'):
    """Count saved reviews toward a comparison-wide target, including the sample."""
    reviewed = [r for r in rows if identity(r) in reviews]
    remaining = max(0, limit - len(reviewed))
    eligible = [r for r in rows if identity(r) not in reviews]

    def priority(row):
        return (*(-int(changes(row, threshold)[s]) for s in SIGNALS[:4]),
                -score_delta(row), identity(row))

    selected = sorted((r for r in eligible if any(changes(r, threshold)[s] for s in signals)),
                      key=priority)
    # Sample truly unchanged review fields, not small changes below the threshold.
    unchanged = [r for r in eligible if not any(changes(r, 0).values())]
    sampled_reviews = sum(not any(changes(r, 0).values()) for r in reviewed)
    sample_size = min(max(0, unchanged_sample - sampled_reviews), len(unchanged), remaining)
    sample = random.Random(42).sample(sorted(unchanged, key=identity), sample_size)
    pending = selected[:remaining - sample_size] + sample
    if status == 'Reviewed':
        return sorted(reviewed, key=priority)
    if status == 'All':
        return pending + sorted(reviewed, key=priority)
    return pending


def load_reviews(path, comparison=None):
    path = Path(path)
    if not path.exists():
        return {}
    reviews = {}
    for row in read_rows(path):
        if not {'review_rating', 'review_note', 'reviewed_at', 'comparison_after'} <= row.keys():
            raise ValueError('Output exists but is not a calibration review workbook. Choose another path.')
        if comparison and any(text(row.get(f)) != text(comparison.get(f))
                              for f in ('comparison_before', 'comparison_after')):
            raise ValueError('Review output belongs to a different comparison. Choose another output path.')
        if row.get('review_rating') not in (1, 2, 3):
            raise ValueError('Saved review has an invalid rating; check the review workbook.')
        key = identity(row)
        if key in reviews:
            raise ValueError(f'Duplicate review rows exist for Job ID {key[1]}; check the file before continuing.')
        for field in SUBGRADE_OPTIONS:
            row[field] = text(row.get(field))  # Older workbooks have no sub-grade columns.
        reviews[key] = row
    if not reviews:
        raise ValueError('Existing output contains no reviews. Choose a new output path.')
    return reviews


def find_review(reviews, job_id):
    """Resolve an ID-only lookup without guessing between source sites."""
    job_id = text(job_id)
    if not job_id:
        raise LookupError('Enter a Job ID to load an existing review.')
    matches = [review for key, review in reviews.items() if key[1] == job_id]
    if not matches:
        raise LookupError(f'No saved review found for Job ID {job_id}.')
    if len(matches) != 1:
        raise ValueError(f'Duplicate review rows exist for Job ID {job_id}; resolve them before editing.')
    return matches[0]


def default_review_path(comparison):
    directory = Path(comparison['comparison_after']).parent
    for path in sorted(directory.glob('calibration_review_*.xlsx'), reverse=True):
        try:
            load_reviews(path, comparison)
            return path
        except ValueError:
            continue
    return directory / f'calibration_review_{date.today().isoformat()}.xlsx'


def save_review(path, row, rating, note='', tags=(), *, decision_quality='',
                score_quality='', requirement_interpretation='', require_existing=False):
    path = Path(path).resolve()
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Review output must end in .xlsx.')
    if any(path == Path(row[f]).resolve() for f in ('comparison_before', 'comparison_after') if row.get(f)):
        raise ValueError('Review output must be separate from the input files.')
    if rating not in (1, 2, 3):
        raise ValueError('Choose a rating: 1, 2, or 3.')
    subgrades = dict(decision_quality=text(decision_quality), score_quality=text(score_quality),
                    requirement_interpretation=text(requirement_interpretation))
    for field, value in subgrades.items():
        if value and value not in SUBGRADE_OPTIONS[field]:
            raise ValueError(f'Invalid {field}: {value}')
    reviews = load_reviews(path, row)
    key = identity(row)
    if require_existing:
        try:
            existing = find_review(reviews, key[1])
        except LookupError as error:
            raise ValueError(str(error)) from error
        if identity(existing) != key:
            raise ValueError('The saved review identity changed. Reload the review before editing.')
    fields = ('source_site', 'job_id', 'title', 'company', 'source', 'comparison_before', 'comparison_after')
    saved = {f: row.get(f, '') for f in fields}
    saved.update({f'{side}_{f}': row.get(f'{side}_{f}', '') for side in ('old', 'new') for f in GRADE_FIELDS})
    for side in ('old', 'new'):
        for field in ('fit_score', 'pre_ai_score', 'learning_score', 'confidence'):
            name = f'{side}_{field}'
            try:
                saved[name] = float(saved[name])
            except (TypeError, ValueError):
                pass  # Keep missing/unavailable values as supplied.
    human_fields = dict(review_rating=rating, **subgrades, review_note=note, issue_tags=';'.join(tags),
                        reviewed=True, reviewed_at=datetime.now(timezone.utc).isoformat(timespec='seconds'))
    saved = dict(reviews.get(key, saved), **human_fields)
    # Keep unrelated columns, formulas, rows and worksheets intact on updates.
    existing_file = path.exists()
    book = load_workbook(path) if existing_file else Workbook()
    sheet = book.active
    headers = [text(cell.value) for cell in sheet[1]] if existing_file else []
    original_column_count = len(headers)
    if not existing_file:
        sheet.title = 'Reviews'
    for field in saved:
        if field not in headers:
            headers.append(field)
            sheet.cell(1, len(headers), field)
    if key in reviews:
        row_number = next(cells[0].row for cells in sheet.iter_rows(min_row=2)
                          if any(cell.value is not None for cell in cells) and
                          identity(dict(zip(headers, (cell.value for cell in cells)))) == key)
        updates = human_fields
    else:
        row_number = sheet.max_row + 1
        updates = saved
    for field, value in updates.items():
        cell = sheet.cell(row_number, headers.index(field) + 1, value)
        if isinstance(value, str):
            cell.data_type = 's'  # Literal notes/IDs, including a leading '='.
    reviews[key] = saved
    for cell in sheet[1][original_column_count:]:
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = 24
    if not existing_file:
        sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    # A locked Excel file leaves the previous complete workbook intact.
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.xlsx', delete=False) as temp:
        temporary = Path(temp.name)
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return reviews
